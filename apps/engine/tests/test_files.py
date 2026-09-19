"""Attached files: the parsers, the limits, and what a turn sees of them."""

import io
import json
import zipfile

import httpx
import pytest

from engine.core.config import Files
from engine.core.doubles import MemorySandbox
from engine.core.types import Attachment, IngestError, SandboxOutput
from engine.memory.ingest import files as reader
from engine.memory.ingest.parsers import (
    PARSERS,
    ParseError,
    archive,
    office,
    parser_for,
    text,
    unhandled,
)


def _docx(paragraphs):
    from docx import Document

    document = Document()
    for line in paragraphs:
        document.add_paragraph(line)
    raw = io.BytesIO()
    document.save(raw)
    return raw.getvalue()


def _pptx(slides):
    from pptx import Presentation

    deck = Presentation()
    for title in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = title
    raw = io.BytesIO()
    deck.save(raw)
    return raw.getvalue()


def _xlsx(rows):
    from openpyxl import Workbook

    book = Workbook()
    for row in rows:
        book.active.append(row)
    raw = io.BytesIO()
    book.save(raw)
    return raw.getvalue()


def _pdf(line):
    """A minimal one-page PDF carrying one line of text, built by hand."""
    content = b"BT /F1 24 Tf 72 700 Td (" + line.encode() + b") Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += str(number).encode() + b" 0 obj\n" + body + b"\nendobj\n"
    start = len(out)
    out += b"xref\n0 " + str(len(objects) + 1).encode() + b"\n0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += (
        b"trailer\n<< /Size "
        + str(len(objects) + 1).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(start).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


def _zip(entries):
    raw = io.BytesIO()
    with zipfile.ZipFile(raw, "w") as archive_file:
        for name, body in entries.items():
            archive_file.writestr(name, body)
    return raw.getvalue()


# ---------------------------------------------------------------- the registry


def test_every_media_type_the_types_module_offers_has_a_parser():
    """A type accepted on the way in and unreadable on the way out is a silent dead end."""
    assert unhandled() == ()


def test_a_type_nothing_reads_has_no_parser():
    assert parser_for("application/x-msdownload") is None
    assert parser_for("application/pdf") is not None


# ---------------------------------------------------------------- the parsers


def test_text_is_read_as_it_is_and_cut_to_the_limit():
    assert text.read(b"hello", 100) == "hello"
    assert text.read(b"hello", 2) == "he"


def test_undecodable_bytes_do_not_fail_the_read():
    assert "hello" in text.read(b"hello \xff\xfe", 100)


def test_json_and_csv_come_back_as_their_own_text():
    body = json.dumps({"budget": 500}).encode()
    assert "budget" in PARSERS["application/json"](body, 100)
    assert "a,b" in PARSERS["text/csv"](b"a,b\n1,2", 100)


def test_a_word_document_comes_back_as_its_paragraphs():
    raw = _docx(["Budget requests", "", "Due the first Friday"])
    read = office.read_docx(raw, 1000)
    assert "Budget requests" in read
    assert "Due the first Friday" in read


def test_a_deck_comes_back_slide_by_slide_under_its_number():
    read = office.read_pptx(_pptx(["Sponsorship", "Timeline"]), 1000)
    assert "Slide 1" in read and "Sponsorship" in read
    assert "Slide 2" in read and "Timeline" in read


def test_a_spreadsheet_comes_back_as_rows_under_its_sheet():
    read = office.read_xlsx(_xlsx([("item", "cost"), ("flyers", 40)]), 1000)
    assert "Sheet" in read
    assert "item\tcost" in read
    assert "flyers\t40" in read


def test_a_pdf_comes_back_as_the_text_on_its_pages():
    read = PARSERS["application/pdf"](_pdf("Budget due 30 September"), 1000)
    assert "Budget due 30 September" in read


def test_a_file_that_is_not_what_it_claims_is_a_parse_error():
    """A parser raises its own error: it runs in the sandbox too, where no engine code is."""
    with pytest.raises(ParseError, match="could not be read"):
        office.read_docx(b"not a docx", 100)
    with pytest.raises(ParseError, match="could not be read"):
        PARSERS["application/pdf"](b"not a pdf", 100)


# ---------------------------------------------------------------- archives


def test_an_archive_reads_the_entries_it_knows_and_names_the_rest():
    raw = _zip({"notes.md": "# Budget", "logo.bin": b"\x00\x01", "sheet.csv": "a,b"})
    read = archive.read(raw, 10_000)
    assert "# Budget" in read
    assert "a,b" in read
    assert "logo.bin" in read and "nothing here reads" in read


def test_a_nested_archive_is_listed_and_not_opened():
    """Opening archives inside archives is how one small upload becomes unbounded work."""
    inner = _zip({"a.txt": "x" * 100})
    read = archive.read(_zip({"inner.zip": inner, "ok.txt": "fine"}), 10_000)
    assert "inner.zip" in read and "not opened" in read
    assert "fine" in read


def test_an_entry_that_unpacks_too_large_is_named_and_skipped():
    raw = _zip({"big.txt": "x" * (archive.MAX_ENTRY_BYTES + 10), "small.txt": "ok"})
    read = archive.read(raw, 10_000_000)
    assert "big.txt" in read and "too large" in read
    assert "ok" in read


def test_an_archive_stops_once_its_entries_unpack_past_the_total():
    many = {f"f{i}.txt": "x" * 4_000_000 for i in range(10)}
    read = archive.read(_zip(many), 100_000_000)
    assert "more than this reads" in read


def test_something_that_is_not_an_archive_is_a_parse_error():
    with pytest.raises(ParseError, match="could not be opened"):
        archive.read(b"not a zip", 100)


# ---------------------------------------------------------------- downloading


def _attachment(media_type="text/plain", name="notes.txt", size=0):
    return Attachment(url="https://cdn/notes.txt", media_type=media_type, name=name, size=size)


def _limits(**over):
    return Files(**over)


def _serving(body=b"", status=200):
    """A transport answering every fetch the same way, and the requests it saw."""
    seen: list[httpx.Request] = []

    def handle(request):
        seen.append(request)
        return httpx.Response(status, content=body)

    return httpx.MockTransport(handle), seen


async def test_a_file_the_platform_calls_too_big_is_refused_before_it_is_fetched():
    transport, seen = _serving(b"x")
    with pytest.raises(IngestError, match="over the"):
        await reader.download(_attachment(size=999), _limits(max_bytes=10), transport)
    assert seen == [], "the size the platform gave is enough to refuse it"


async def test_a_file_larger_than_it_claimed_is_still_refused():
    transport, _ = _serving(b"x" * 50)
    with pytest.raises(IngestError, match="over the"):
        await reader.download(_attachment(), _limits(max_bytes=10), transport)


async def test_a_fetch_that_fails_says_so_without_the_url():
    transport, _ = _serving(status=404)
    with pytest.raises(IngestError, match="HTTP 404"):
        await reader.download(_attachment(), _limits(), transport)


# ---------------------------------------------------------------- reading into a document


async def test_a_read_file_becomes_a_document_named_by_its_bytes():
    transport, _ = _serving(b"the budget")

    document = await reader.read(_attachment(), _limits(), transport)

    assert document.source == reader.SOURCE
    assert document.source_id == reader.source_id(b"the budget")
    assert document.text == "the budget"
    assert document.title == "notes.txt"
    assert document.metadata["media_type"] == "text/plain"


async def test_the_same_bytes_are_the_same_document():
    assert reader.source_id(b"same") == reader.source_id(b"same")
    assert reader.source_id(b"same") != reader.source_id(b"other")


async def test_a_type_with_no_parser_says_so_rather_than_fetching():
    transport, seen = _serving(b"x")
    with pytest.raises(IngestError, match="nothing here reads"):
        await reader.read(_attachment(media_type="image/png"), _limits(), transport)
    assert seen == []


# ---------------------------------------------------------------- the turn


def _manager(cfg, documents, files=None):
    from engine.core.doubles import (
        FixedEmbedder,
        MemoryCollaboration,
        MemoryConversation,
        MemoryOrgContext,
    )
    from engine.memory.manager import MemoryManager

    return MemoryManager(
        memory=cfg.memory,
        conversation=MemoryConversation(),
        org_context=MemoryOrgContext(),
        embedder=FixedEmbedder([0.1]),
        documents=documents,
        collaboration=MemoryCollaboration(),
        settings=cfg.collaboration,
        files=files or Files(),
    )


async def test_a_turn_with_no_files_reads_nothing(cfg, ctx):
    from engine.core.doubles import MemoryDocuments

    documents = MemoryDocuments()
    assert await _manager(cfg, documents).read_files(ctx) == ""
    assert documents.chunks == []


async def test_files_off_reads_nothing_even_when_one_is_attached(cfg, ctx):
    from dataclasses import replace

    from engine.core.doubles import MemoryDocuments

    documents = MemoryDocuments()
    with_file = replace(ctx, files=(_attachment(),))
    manager = _manager(cfg, documents, Files(enabled=False))

    assert await manager.read_files(with_file) == ""
    assert documents.chunks == []


async def test_a_file_that_cannot_be_read_says_so_and_does_not_fail_the_turn(cfg, ctx):
    """The rest of the message is still worth answering."""
    from dataclasses import replace

    from engine.core.doubles import MemoryDocuments

    documents = MemoryDocuments()
    broken = Attachment(url="https://cdn/x.pdf", media_type="application/pdf", name="budget.pdf")
    seen = await _manager(cfg, documents).read_files(replace(ctx, files=(broken,)))

    assert "budget.pdf" in seen
    assert documents.chunks == []


async def test_only_max_per_message_files_are_read(cfg, ctx):
    from dataclasses import replace

    from engine.core.doubles import MemoryDocuments

    many = tuple(
        Attachment(url=f"https://cdn/{i}.png", media_type="image/png", name=f"{i}.png")
        for i in range(10)
    )
    seen = await _manager(cfg, MemoryDocuments(), Files(max_per_message=2)).read_files(
        replace(ctx, files=many)
    )
    assert len(seen.splitlines()) == 3, "two files, one blank line between them"


# ---------------------------------------------------------------- parsing in the sandbox


def sandboxed(cfg, **over):
    """The file settings with sandbox parsing on."""
    return cfg.files.model_copy(update={"sandbox": True, **over})


def _said(text):
    """What the extractor prints for a file it read."""
    return SandboxOutput(exit_code=0, stdout=json.dumps({"text": text}), stderr="", session="files")


@pytest.mark.anyio
async def test_with_the_sandbox_on_the_bytes_are_parsed_in_the_container(cfg, ctx):
    transport, _ = _serving(b"budget notes")
    box = MemorySandbox(
        default=SandboxOutput(
            exit_code=0, stdout='{"text": "budget notes"}', stderr="", session="files"
        )
    )
    attachment = Attachment(url="https://x/notes.txt", media_type="text/plain", name="notes.txt")

    document = await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=box)

    assert document.text == "budget notes"
    # The bytes went into the workspace, and the command named the extractor.
    assert any(key.endswith(".txt") for key in box.written)
    assert "/opt/zipy/extract.py" in box.ran[0][1]
    assert box.ran[0][2] == reader.PARSE_SESSION


@pytest.mark.anyio
async def test_the_sandbox_never_sees_another_orgs_workspace(cfg, ctx):
    transport, _ = _serving(b"hi")
    box = MemorySandbox(default=_said("hi"))
    attachment = Attachment(url="https://x/a.txt", media_type="text/plain", name="a.txt")

    await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=box)

    assert all(key.startswith(f"{ctx.org_id}/") for key in box.written)
    assert box.ran[0][0] == ctx.org_id


@pytest.mark.anyio
async def test_the_sandbox_is_never_asked_to_run_what_a_file_is_called(cfg, ctx):
    """A name is not a command: it reaches the container quoted, and not as the stored name."""
    transport, _ = _serving(b"hi")
    box = MemorySandbox(default=_said("hi"))
    attachment = Attachment(url="https://x/a.txt", media_type="text/plain", name="; rm -rf / #.txt")

    await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=box)

    assert "rm -rf" not in box.ran[0][1]


@pytest.mark.anyio
async def test_a_file_the_sandbox_could_not_read_says_so_rather_than_answering(cfg, ctx):
    transport, _ = _serving(b"not a pdf")
    box = MemorySandbox()
    box.answers["x"] = SandboxOutput(exit_code=0, stdout="", stderr="", session="files")
    attachment = Attachment(url="https://x/a.pdf", media_type="application/pdf", name="a.pdf")

    with pytest.raises(IngestError, match="could not be read"):
        await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=box)


@pytest.mark.anyio
async def test_an_extractor_error_reaches_the_person_as_the_reason(cfg, ctx):
    transport, _ = _serving(b"x")
    box = MemorySandbox()
    box.answers = {}
    attachment = Attachment(url="https://x/a.pdf", media_type="application/pdf", name="a.pdf")

    async def failed(request_ctx, request):
        box.ran.append((request_ctx.org_id, request.command, request.session))
        return SandboxOutput(
            exit_code=1, stdout='{"error": "that file is not a pdf"}', stderr="", session="files"
        )

    box.run = failed  # type: ignore[method-assign]

    with pytest.raises(IngestError, match="not a pdf"):
        await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=box)


@pytest.mark.anyio
async def test_sandbox_parsing_never_quietly_falls_back_to_this_process(cfg, ctx):
    """The point is that an untrusted file is not read beside the org's credentials."""
    transport, _ = _serving(b"plain text")
    attachment = Attachment(url="https://x/a.txt", media_type="text/plain", name="a.txt")

    with pytest.raises(IngestError, match="no sandbox is wired"):
        await reader.read(attachment, sandboxed(cfg), transport, ctx=ctx, sandbox=None)


def test_the_sandbox_is_off_until_someone_turns_it_on(cfg):
    assert cfg.files.sandbox is False
